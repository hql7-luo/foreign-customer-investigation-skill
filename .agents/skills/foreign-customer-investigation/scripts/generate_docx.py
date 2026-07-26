"""Generate a compact, business-style Word investigation report."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

try:
    from .file_utils import (
        ensure_output_dir,
        evidence_label,
        join_items,
        localized,
        output_basename,
    )
    from .generate_json import load_report
    from .models import EvidenceStatus, InvestigationReport, ReportLanguage, ReportMode
except ImportError:  # pragma: no cover
    from file_utils import ensure_output_dir, evidence_label, join_items, localized, output_basename
    from generate_json import load_report
    from models import EvidenceStatus, InvestigationReport, ReportLanguage, ReportMode


# Resolved "compact_reference_guide" tokens with a named "two_page_report" override.
PAGE_WIDTH_IN = 8.5
PAGE_HEIGHT_IN = 11.0
MARGIN_IN = 0.62
CONTENT_WIDTH_IN = PAGE_WIDTH_IN - 2 * MARGIN_IN
BODY_FONT = "Calibri"
CJK_FONT = "Hiragino Sans GB"
BODY_SIZE_PT = 9.0
HEADING_SIZE_PT = 11.5
TABLE_SIZE_PT = 8.1
DEEP_BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
BORDER_BLUE_GRAY = "8EA9C1"
MUTED = "5B6573"


def _set_run_font(run, size: float, bold: bool = False, color: str | None = None) -> None:
    font_name = CJK_FONT if any(ord(character) > 127 for character in run.text) else BODY_FONT
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    run_properties = run._element.get_or_add_rPr()
    fonts = run_properties.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run_properties.insert(0, fonts)
    fonts.set(qn("w:ascii"), font_name)
    fonts.set(qn("w:hAnsi"), font_name)
    fonts.set(qn("w:eastAsia"), CJK_FONT)
    fonts.set(qn("w:cs"), font_name)


def _configure_document(document: Document) -> None:
    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(PAGE_WIDTH_IN)
    section.page_height = Inches(PAGE_HEIGHT_IN)
    section.top_margin = Inches(MARGIN_IN)
    section.bottom_margin = Inches(MARGIN_IN)
    section.left_margin = Inches(MARGIN_IN)
    section.right_margin = Inches(MARGIN_IN)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(BODY_SIZE_PT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(2.5)
    normal.paragraph_format.line_spacing = 1.05

    heading = document.styles["Heading 1"]
    heading.font.name = BODY_FONT
    heading.font.size = Pt(HEADING_SIZE_PT)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor.from_string(DEEP_BLUE)
    heading._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    heading.paragraph_format.space_before = Pt(5)
    heading.paragraph_format.space_after = Pt(2.5)
    heading.paragraph_format.keep_with_next = True

    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.comments = ""
    document.core_properties.keywords = ""


def _shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_margins(cell, top: int = 70, start: int = 100, bottom: int = 70, end: int = 100) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_geometry(table, widths_in: list[float]) -> None:
    total_dxa = round(sum(widths_in) * 1440)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table_properties = table._tbl.tblPr

    table_width = table_properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table_properties.append(table_width)
    table_width.set(qn("w:w"), str(total_dxa))
    table_width.set(qn("w:type"), "dxa")

    indent = table_properties.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        table_properties.append(indent)
    indent.set(qn("w:w"), "0")
    indent.set(qn("w:type"), "dxa")

    layout = table_properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        table_properties.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    widths_dxa = [round(width * 1440) for width in widths_in]
    for width in widths_dxa:
        grid_column = OxmlElement("w:gridCol")
        grid_column.set(qn("w:w"), str(width))
        grid.append(grid_column)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[min(index, len(widths_dxa) - 1)]
            cell.width = Inches(widths_in[min(index, len(widths_in) - 1)])
            properties = cell._tc.get_or_add_tcPr()
            cell_width = properties.find(qn("w:tcW"))
            if cell_width is None:
                cell_width = OxmlElement("w:tcW")
                properties.append(cell_width)
            cell_width.set(qn("w:w"), str(width))
            cell_width.set(qn("w:type"), "dxa")
            _set_cell_margins(cell)


def _set_table_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), BORDER_BLUE_GRAY)


def _repeat_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def _set_cell_text(
    cell,
    text: str,
    *,
    bold: bool = False,
    color: str | None = None,
    align: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.LEFT,
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    _set_run_font(run, TABLE_SIZE_PT, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _format_table(table, widths_in: list[float], header: bool = True) -> None:
    _set_table_geometry(table, widths_in)
    _set_table_borders(table)
    if header:
        _repeat_header(table.rows[0])
        for cell in table.rows[0].cells:
            _shade(cell, LIGHT_BLUE)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    _set_run_font(run, TABLE_SIZE_PT, bold=True, color=DEEP_BLUE)


def _add_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(text, style="Heading 1")
    for run in paragraph.runs:
        _set_run_font(run, HEADING_SIZE_PT, bold=True, color=DEEP_BLUE)


def _add_title(document: Document, report: InvestigationReport) -> None:
    title_suffix = localized(
        report.report_language,
        "客户背景调查报告",
        "Customer Background Investigation Report",
    )
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(1)
    run = paragraph.add_run(f"{report.customer.legal_name}\n{title_suffix}")
    _set_run_font(run, 15.5, bold=True, color=DEEP_BLUE)

    date_label = localized(report.report_language, "调查日期", "Investigation date")
    date_paragraph = document.add_paragraph()
    date_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_paragraph.paragraph_format.space_after = Pt(3)
    date_run = date_paragraph.add_run(f"{date_label}: {report.investigation_date.isoformat()}")
    _set_run_font(date_run, 8.5, color=MUTED)

    if report.fictional_notice:
        notice = document.add_paragraph()
        notice.alignment = WD_ALIGN_PARAGRAPH.CENTER
        notice.paragraph_format.space_after = Pt(2)
        notice_run = notice.add_run(report.fictional_notice)
        _set_run_font(notice_run, 8, bold=True, color="9B1C1C")


def _add_full_report(document: Document, report: InvestigationReport) -> None:
    language = report.report_language
    _add_heading(document, localized(language, "1. 结论", "1. Conclusion"))
    conclusion = document.add_paragraph(report.conclusion)
    conclusion.paragraph_format.space_after = Pt(3)
    for run in conclusion.runs:
        _set_run_font(run, BODY_SIZE_PT)

    _add_heading(
        document,
        localized(language, "2. 九项背景调查总表", "2. Nine-part investigation summary"),
    )
    table = document.add_table(rows=1, cols=3)
    headers = [
        localized(language, "序号", "No."),
        localized(language, "调查项目", "Investigation item"),
        localized(language, "核心发现", "Core finding"),
    ]
    for cell, header in zip(table.rows[0].cells, headers, strict=True):
        _set_cell_text(cell, header, bold=True, color=DEEP_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER)
    for item in sorted(report.investigation_items, key=lambda value: value.number):
        cells = table.add_row().cells
        _set_cell_text(cells[0], str(item.number), align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_cell_text(cells[1], item.item)
        finding = f"{evidence_label(item.evidence_status, language)} {item.core_finding}"
        _set_cell_text(cells[2], finding)
    _format_table(table, [0.48, 2.0, CONTENT_WIDTH_IN - 2.48])

    _add_heading(document, localized(language, "3. 业务员行动卡", "3. Sales action card"))
    card = document.add_table(rows=0, cols=2)
    products = join_items(report.recommended_products, language)
    grade_score = f"{report.scoring.final_grade.value} / {report.scoring.total_score}"
    gaps = join_items(report.evidence_gaps, language)
    rows = [
        (localized(language, "最终等级", "Final grade"), grade_score),
        (localized(language, "首推产品", "Priority products"), products),
        (localized(language, "首次联系重点", "First-contact focus"), report.first_contact_focus),
        (localized(language, "当前证据缺口", "Current evidence gaps"), gaps),
        (localized(language, "下一步动作", "Priority action"), report.priority_action),
        (
            localized(language, "暂不建议动作", "Action not recommended"),
            report.not_recommended_action,
        ),
    ]
    for label, value in rows:
        cells = card.add_row().cells
        _set_cell_text(cells[0], label, bold=True, color=DEEP_BLUE)
        _shade(cells[0], LIGHT_BLUE)
        _set_cell_text(cells[1], value)
    _format_table(card, [1.55, CONTENT_WIDTH_IN - 1.55], header=False)

    _add_heading(document, localized(language, "4. 核心来源", "4. Core sources"))
    for source in report.core_sources:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0.5)
        run = paragraph.add_run(source.url)
        _set_run_font(run, 7.8, color=MUTED)


def _add_simplified_report(document: Document, report: InvestigationReport) -> None:
    language = report.report_language
    _add_heading(document, localized(language, "1. 结论", "1. Conclusion"))
    document.add_paragraph(report.conclusion)
    for run in document.paragraphs[-1].runs:
        _set_run_font(run, BODY_SIZE_PT)

    _add_heading(document, localized(language, "2. 开发价值判断", "2. Development value"))
    document.add_paragraph(report.development_value)
    for run in document.paragraphs[-1].runs:
        _set_run_font(run, BODY_SIZE_PT)

    _add_heading(
        document,
        localized(language, "3. 已确认信息清单", "3. Confirmed information"),
    )
    confirmed = [
        item
        for item in report.investigation_items
        if item.evidence_status in {EvidenceStatus.CONFIRMED, EvidenceStatus.PARTIALLY_CONFIRMED}
    ]
    if not confirmed:
        paragraph = document.add_paragraph(
            localized(language, "未找到可靠公开信息", "No reliable public information was found")
        )
        for run in paragraph.runs:
            _set_run_font(run, BODY_SIZE_PT)
    for item in confirmed:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.first_line_indent = Inches(-0.15)
        paragraph.paragraph_format.space_after = Pt(1)
        run = paragraph.add_run(
            f"{evidence_label(item.evidence_status, language)} {item.core_finding}"
        )
        _set_run_font(run, BODY_SIZE_PT)

    _add_heading(
        document,
        localized(language, "4. 建议业务员直接核实的问题", "4. Questions to verify directly"),
    )
    for index, question in enumerate(
        sorted(report.follow_up_questions, key=lambda value: value.priority), start=1
    ):
        paragraph = document.add_paragraph(f"{index}. {question.question}")
        for run in paragraph.runs:
            _set_run_font(run, BODY_SIZE_PT)

    _add_heading(document, localized(language, "5. 核心来源", "5. Core sources"))
    for source in report.core_sources:
        paragraph = document.add_paragraph(source.url)
        paragraph.paragraph_format.space_after = Pt(0.5)
        for run in paragraph.runs:
            _set_run_font(run, 7.8, color=MUTED)


def generate_docx(report: InvestigationReport, output_path: str | Path) -> Path:
    """Create one report and scrub common author metadata."""

    destination = Path(output_path).expanduser().resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    _configure_document(document)
    _add_title(document, report)
    if report.mode == ReportMode.SIMPLIFIED:
        _add_simplified_report(document, report)
    else:
        _add_full_report(document, report)
    document.save(destination)
    return destination


def default_filename(report: InvestigationReport) -> str:
    base = output_basename(report)
    if report.report_language == ReportLanguage.ZH_CN:
        suffix = "客户背景调查报告"
    elif report.report_language == ReportLanguage.EN:
        suffix = "Customer_Background_Report"
    else:
        suffix = "客户背景调查报告_Customer_Background_Report"
    return f"{base}_{suffix}.docx"


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input_json", help="Validated investigation data")
    parser.add_argument("--output-dir", default="outputs")
    args = parser.parse_args()
    report = load_report(args.input_json)
    directory = ensure_output_dir(args.output_dir)
    print(generate_docx(report, directory / default_filename(report)))


if __name__ == "__main__":
    main()
