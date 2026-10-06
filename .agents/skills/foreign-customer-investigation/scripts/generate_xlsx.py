"""Generate a styled, formula-auditable Excel investigation workbook."""

from __future__ import annotations

import argparse
from datetime import date
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

try:
    from .file_utils import (
        ensure_output_dir,
        evidence_label,
        join_items,
        localized,
        output_basename,
    )
    from .generate_json import load_report
    from .models import (
        Grade,
        InvestigationReport,
        RankingEntry,
        ReportLanguage,
        ReportMode,
        rank_reports,
    )
except ImportError:  # pragma: no cover
    from file_utils import ensure_output_dir, evidence_label, join_items, localized, output_basename
    from generate_json import load_report
    from models import Grade, InvestigationReport, RankingEntry, ReportLanguage, ReportMode, rank_reports


NAVY = "1F4E78"
LIGHT_BLUE = "D9EAF7"
VERY_LIGHT_BLUE = "EDF5FB"
WHITE = "FFFFFF"
TEXT = "263238"
GRID = "B8C6D1"
GREEN = "C6EFCE"
YELLOW = "FFEB9C"
ORANGE = "FCE4D6"
RED = "FFC7CE"
THIN = Side(style="thin", color=GRID)


SHEET_NAMES = {
    ReportLanguage.ZH_CN: [
        "客户概览",
        "九项调查",
        "评分明细",
        "证据与来源",
        "跟进问题",
        "开发排名",
    ],
    ReportLanguage.EN: [
        "Customer Overview",
        "Investigation",
        "Scoring Details",
        "Evidence and Sources",
        "Follow-up Questions",
        "Development Ranking",
    ],
    ReportLanguage.BILINGUAL: [
        "客户概览 Overview",
        "九项调查 Investigation",
        "评分明细 Scoring",
        "证据来源 Evidence",
        "跟进问题 Follow-up",
        "开发排名 Ranking",
    ],
}


def _append_literal_row(sheet: Worksheet, values: Iterable[object]) -> None:
    """Keep research text literal; only the scoring builder may create formulas."""

    row_values = list(values)
    sheet.append(row_values)
    for column, value in enumerate(row_values, start=1):
        if isinstance(value, str):
            # openpyxl infers formulas/errors from strings such as =1+1 or #N/A.
            # Explicit string cells preserve the original text without executing it.
            sheet.cell(sheet.max_row, column).data_type = "s"


def _headers(language: ReportLanguage, zh: list[str], en: list[str]) -> list[str]:
    return [localized(language, left, right) for left, right in zip(zh, en, strict=True)]


def _style_header(sheet: Worksheet, row: int = 1) -> None:
    for cell in sheet[row]:
        if cell.value is None:
            continue
        cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        cell.font = Font(name="Calibri", size=10, bold=True, color=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=NAVY))
    sheet.row_dimensions[row].height = 28


def _style_body(sheet: Worksheet) -> None:
    for row in sheet.iter_rows():
        for cell in row:
            if cell.row == 1:
                continue
            cell.font = Font(name="Calibri", size=9, color=TEXT)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.value is not None:
                cell.border = Border(bottom=THIN)


def _finalize_sheet(
    sheet: Worksheet,
    widths: list[float],
    *,
    filter_range: str | None = None,
    freeze: str = "A2",
    landscape: bool = False,
) -> None:
    _style_header(sheet)
    _style_body(sheet)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = freeze
    if filter_range:
        sheet.auto_filter.ref = filter_range
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.orientation = "landscape" if landscape else "portrait"
    sheet.page_setup.paperSize = "9"  # ISO A4
    sheet.sheet_view.zoomScale = 90


def _add_grade_formatting(sheet: Worksheet, cell_range: str) -> None:
    for grade, fill in (
        ("A", GREEN),
        ("B", YELLOW),
        ("C", ORANGE),
        ("D", RED),
    ):
        sheet.conditional_formatting.add(
            cell_range,
            FormulaRule(
                formula=[f'{cell_range.split(":")[0]}="{grade}"'],
                font=Font(name="Calibri", size=9, bold=True, color=TEXT),
                fill=PatternFill("solid", fgColor=fill),
            ),
        )


def _overview_sheet(sheet: Worksheet, report: InvestigationReport) -> None:
    language = report.report_language
    _append_literal_row(sheet, _headers(language, ["字段", "内容"], ["Field", "Value"]))
    rows: list[tuple[str, object]] = [
        (
            localized(language, "示例声明", "Example notice"),
            report.fictional_notice,
        ),
        (localized(language, "公司名称", "Company name"), report.customer.legal_name),
        (localized(language, "品牌", "Brand"), report.customer.brand or report.customer.trading_name),
        (localized(language, "国家", "Country"), report.customer.country_name or report.customer.country_code),
        (localized(language, "联系人", "Contact"), report.customer.contact.name),
        (localized(language, "职位", "Job title"), report.customer.contact.job_title),
        (localized(language, "官网", "Website"), report.customer.website),
        (localized(language, "邮箱", "Email"), report.customer.contact.email),
        (localized(language, "电话", "Phone"), report.customer.contact.phone),
        (localized(language, "主营业务", "Main business"), report.business_analysis.main_business),
        (localized(language, "最终等级", "Final grade"), report.scoring.final_grade.value),
        (localized(language, "最终得分", "Final score"), report.scoring.total_score),
        (
            localized(language, "首推产品", "Priority products"),
            join_items(report.recommended_products, language),
        ),
        (localized(language, "最优先行动", "Priority action"), report.priority_action),
        (localized(language, "最大风险", "Maximum risk"), report.maximum_risk),
        (localized(language, "调查日期", "Investigation date"), report.investigation_date),
    ]
    for label, value in rows:
        _append_literal_row(sheet, [label, value])
    for cell in sheet["A"][1:]:
        cell.fill = PatternFill("solid", fgColor=VERY_LIGHT_BLUE)
        cell.font = Font(name="Calibri", size=9, bold=True, color=NAVY)
    grade_row = next(row for row in range(2, sheet.max_row + 1) if "grade" in str(sheet.cell(row, 1).value).lower() or "等级" in str(sheet.cell(row, 1).value))
    _add_grade_formatting(sheet, f"B{grade_row}")
    for row in range(2, sheet.max_row + 1):
        if isinstance(sheet.cell(row, 2).value, date):
            sheet.cell(row, 2).number_format = "yyyy-mm-dd"
    _finalize_sheet(sheet, [24, 88], filter_range=f"A1:B{sheet.max_row}")


def _investigation_sheet(sheet: Worksheet, report: InvestigationReport) -> None:
    language = report.report_language
    headers = _headers(
        language,
        ["序号", "调查项目", "核心发现", "证据状态", "信息来源", "是否需要客户确认"],
        ["No.", "Investigation item", "Core finding", "Evidence status", "Sources", "Customer confirmation needed"],
    )
    _append_literal_row(sheet, headers)
    for item in sorted(report.investigation_items, key=lambda value: value.number):
        _append_literal_row(
            sheet,
            [
                item.number,
                item.item,
                item.core_finding,
                evidence_label(item.evidence_status, language),
                ", ".join(item.source_ids),
                localized(language, "是" if item.needs_customer_confirmation else "否", "Yes" if item.needs_customer_confirmation else "No"),
            ]
        )
    _finalize_sheet(
        sheet,
        [8, 31, 68, 24, 20, 25],
        filter_range=f"A1:F{sheet.max_row}",
        landscape=True,
    )


def _scoring_sheet(sheet: Worksheet, report: InvestigationReport) -> None:
    language = report.report_language
    _append_literal_row(
        sheet,
        _headers(
            language,
            ["评分维度", "评分项目", "满分", "实际得分", "评分依据", "证据状态"],
            ["Scoring dimension", "Scoring item", "Maximum", "Actual score", "Rationale", "Evidence status"],
        )
    )
    dimensions = [
        report.scoring.authenticity,
        report.scoring.business_fit,
        report.scoring.purchase_signal,
        report.scoring.purchase_potential,
        report.scoring.contact_execution,
    ]
    for dimension in dimensions:
        for item in dimension.items:
            _append_literal_row(
                sheet,
                [
                    dimension.name,
                    item.name,
                    item.max_score,
                    item.score,
                    item.rationale,
                    evidence_label(item.evidence_status, language),
                ]
            )
    for item in report.scoring.risk_deduction.items:
        _append_literal_row(
            sheet,
            [
                report.scoring.risk_deduction.name,
                item.name,
                "0 to -10",
                item.score,
                item.rationale,
                evidence_label(item.evidence_status, language),
            ]
        )
    detail_end = sheet.max_row
    labels: dict[str, int] = {}

    def add_check(key: str, label_zh: str, label_en: str, value: object) -> int:
        sheet.append([localized(language, label_zh, label_en), "", "", value, "", ""])
        labels[key] = sheet.max_row
        return sheet.max_row

    calculated_total_row = add_check(
        "calculated_total",
        "公式合计",
        "Formula total",
        f"=MAX(0,MIN(100,SUM(D2:D{detail_end})))",
    )
    model_total_row = add_check("model_total", "模型总分", "Model total", report.scoring.total_score)
    add_check(
        "total_validation",
        "总分校验",
        "Total validation",
        f'=IF(D{calculated_total_row}=D{model_total_row},"PASS","FAIL")',
    )
    add_check(
        "explicit",
        "明确采购信号",
        "Explicit purchase signal",
        "Yes" if report.explicit_purchase_signal else "No",
    )
    add_check(
        "doubtful",
        "主体真实性存疑",
        "Authenticity doubtful",
        "Yes" if report.verification.authenticity_doubtful else "No",
    )
    add_check(
        "consistent",
        "信息基本一致",
        "Information consistent",
        "Yes" if report.verification.information_consistent else "No",
    )
    add_check(
        "registry",
        "工商主体已确认",
        "Registry confirmed",
        "Yes" if report.verification.business_registry_confirmed else "No",
    )
    both_unconfirmed = not report.customer.website and not report.verification.business_registry_confirmed
    add_check(
        "website_registry",
        "官网且工商均未确认",
        "Website and registry unconfirmed",
        "Yes" if both_unconfirmed else "No",
    )
    add_check(
        "simplified",
        "简化模式",
        "Simplified mode",
        "Yes" if report.mode == ReportMode.SIMPLIFIED else "No",
    )
    severe = any(
        risk.deduction <= -8
        and risk.category.lower() in {"sanctions", "compliance", "fraud", "entity", "payment"}
        for risk in report.risks
    )
    add_check("severe", "严重风险", "Severe risk", "Yes" if severe else "No")
    grade_formula = (
        f'=IF(D{labels["severe"]}="Yes","D",'
        f'IF(D{calculated_total_row}<40,"D",'
        f'IF(D{labels["simplified"]}="Yes","C",'
        f'IF(D{labels["website_registry"]}="Yes","C",'
        f'IF(D{calculated_total_row}<60,"C",'
        f'IF(D{calculated_total_row}<80,"B",'
        f'IF(AND(D{labels["explicit"]}="Yes",D{labels["doubtful"]}="No",'
        f'D{labels["consistent"]}="Yes",D{labels["registry"]}="Yes"),"A","B")))))))'
    )
    expected_grade_row = add_check("expected_grade", "公式等级", "Formula grade", grade_formula)
    model_grade_row = add_check(
        "model_grade", "模型等级", "Model grade", report.scoring.final_grade.value
    )
    add_check(
        "grade_validation",
        "等级校验",
        "Grade validation",
        f'=IF(D{expected_grade_row}=D{model_grade_row},"PASS","FAIL")',
    )
    for row in range(detail_end + 1, sheet.max_row + 1):
        sheet.cell(row, 1).fill = PatternFill("solid", fgColor=VERY_LIGHT_BLUE)
        sheet.cell(row, 1).font = Font(name="Calibri", size=9, bold=True, color=NAVY)
        sheet.cell(row, 4).font = Font(name="Calibri", size=9, bold=True, color=TEXT)
    sheet.conditional_formatting.add(
        f"D2:D{detail_end}",
        ColorScaleRule(
            start_type="min",
            start_color=RED,
            mid_type="percentile",
            mid_value=50,
            mid_color=YELLOW,
            end_type="max",
            end_color=GREEN,
        ),
    )
    _add_grade_formatting(sheet, f"D{model_grade_row}")
    _finalize_sheet(
        sheet,
        [25, 32, 12, 14, 68, 25],
        filter_range=f"A1:F{detail_end}",
        landscape=True,
    )


def _evidence_sheet(sheet: Worksheet, report: InvestigationReport) -> None:
    language = report.report_language
    _append_literal_row(
        sheet,
        _headers(
            language,
            ["来源编号", "页面标题", "完整网址", "来源类型", "访问日期", "支持的结论", "可信度", "信息年份"],
            ["Source ID", "Page title", "Full URL", "Source type", "Access date", "Supported conclusion", "Credibility", "Information year"],
        )
    )
    for source in report.sources:
        _append_literal_row(
            sheet,
            [
                source.source_id,
                source.page_title,
                source.url,
                source.source_type,
                source.access_date,
                source.supports,
                source.credibility,
                source.information_year,
            ]
        )
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, 5).number_format = "yyyy-mm-dd"
    _finalize_sheet(
        sheet,
        [13, 32, 60, 20, 15, 55, 13, 15],
        filter_range=f"A1:H{sheet.max_row}",
        landscape=True,
    )


def _questions_sheet(sheet: Worksheet, report: InvestigationReport) -> None:
    language = report.report_language
    _append_literal_row(
        sheet,
        _headers(
            language,
            ["优先级", "需要确认的问题", "确认原因", "客户回答", "后续动作"],
            ["Priority", "Question to confirm", "Reason", "Customer answer", "Subsequent action"],
        )
    )
    for question in sorted(report.follow_up_questions, key=lambda value: value.priority):
        _append_literal_row(
            sheet,
            [
                question.priority,
                question.question,
                question.reason,
                question.customer_answer,
                question.subsequent_action,
            ]
        )
    _finalize_sheet(
        sheet,
        [12, 55, 45, 35, 40],
        filter_range=f"A1:E{sheet.max_row}",
        landscape=True,
    )


def _ranking_sheet(
    sheet: Worksheet,
    report: InvestigationReport,
    ranking: Iterable[RankingEntry],
) -> None:
    language = report.report_language
    _append_literal_row(
        sheet,
        _headers(
            language,
            ["排名", "公司名称", "最终等级", "最终得分", "首推产品", "核心采购信号", "最大风险", "下一步动作"],
            ["Rank", "Company name", "Final grade", "Final score", "Priority products", "Core purchase signal", "Maximum risk", "Next action"],
        )
    )
    for entry in ranking:
        _append_literal_row(
            sheet,
            [
                entry.rank,
                entry.company_name,
                entry.final_grade.value,
                entry.total_score,
                join_items(entry.recommended_products, language),
                entry.core_purchase_signal,
                entry.maximum_risk,
                entry.priority_action,
            ]
        )
    _add_grade_formatting(sheet, f"C2:C{max(2, sheet.max_row)}")
    sheet.conditional_formatting.add(
        f"D2:D{max(2, sheet.max_row)}",
        ColorScaleRule(
            start_type="min",
            start_color=RED,
            mid_type="percentile",
            mid_value=50,
            mid_color=YELLOW,
            end_type="max",
            end_color=GREEN,
        ),
    )
    _finalize_sheet(
        sheet,
        [9, 35, 13, 13, 38, 48, 48, 50],
        filter_range=f"A1:H{sheet.max_row}",
        landscape=True,
    )


def generate_xlsx(
    report: InvestigationReport,
    output_path: str | Path,
    ranking: list[RankingEntry] | None = None,
) -> Path:
    """Create all required worksheets from one validated report object."""

    destination = Path(output_path).expanduser().resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = ""
    workbook.properties.lastModifiedBy = ""
    workbook.properties.title = "Customer investigation data"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    names = SHEET_NAMES[report.report_language]
    sheets = [workbook.create_sheet(name) for name in names]
    _overview_sheet(sheets[0], report)
    _investigation_sheet(sheets[1], report)
    _scoring_sheet(sheets[2], report)
    _evidence_sheet(sheets[3], report)
    _questions_sheet(sheets[4], report)
    _ranking_sheet(sheets[5], report, ranking or rank_reports([report]))
    workbook.save(destination)
    return destination


def default_filename(report: InvestigationReport) -> str:
    base = output_basename(report)
    if report.report_language == ReportLanguage.EN:
        suffix = "Customer_Investigation_Data"
    elif report.report_language == ReportLanguage.BILINGUAL:
        suffix = "客户背景调查数据_Customer_Investigation_Data"
    else:
        suffix = "客户背景调查数据"
    return f"{base}_{suffix}.xlsx"


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input_json", help="Validated investigation data")
    parser.add_argument("--output-dir", default="outputs")
    args = parser.parse_args()
    report = load_report(args.input_json)
    directory = ensure_output_dir(args.output_dir)
    print(generate_xlsx(report, directory / default_filename(report)))


if __name__ == "__main__":
    main()
