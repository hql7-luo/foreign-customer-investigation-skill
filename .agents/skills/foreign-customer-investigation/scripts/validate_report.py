"""Validate model rules and cross-check Word, Excel, and JSON conclusions."""

from __future__ import annotations

import argparse
import json
import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook

try:
    from .file_utils import join_items, localized
    from .generate_json import load_report
    from .generate_xlsx import SHEET_NAMES
    from .models import (
        Grade,
        InvestigationReport,
        RankingEntry,
        ReportMode,
        derive_grade,
        should_use_simplified_mode,
    )
except ImportError:  # pragma: no cover
    from file_utils import join_items, localized
    from generate_json import load_report
    from generate_xlsx import SHEET_NAMES
    from models import Grade, InvestigationReport, RankingEntry, ReportMode, derive_grade, should_use_simplified_mode


PLACEHOLDER_PATTERNS = [
    re.compile(pattern, re.IGNORECASE)
    for pattern in (
        r"\bTODO\b",
        r"\bTBD\b",
        r"\bPLACEHOLDER\b",
        r"\{\{.+?\}\}",
        r"<placeholder>",
        r"待填写",
    )
]
FORBIDDEN_OUTPUT_TERMS = [
    "搜索过程",
    "内部推理",
    "评分草稿",
    "自查清单",
    "search process",
    "internal reasoning",
    "chain of thought",
    "scoring draft",
    "self-check checklist",
]


@dataclass(slots=True)
class ValidationResult:
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def passed(self) -> bool:
        return not self.errors

    def extend(self, other: "ValidationResult") -> None:
        self.errors.extend(other.errors)
        self.warnings.extend(other.warnings)


def _flatten_strings(value: Any) -> list[str]:
    if isinstance(value, str):
        return [value]
    if isinstance(value, dict):
        flattened: list[str] = []
        for child in value.values():
            flattened.extend(_flatten_strings(child))
        return flattened
    if isinstance(value, list):
        flattened = []
        for child in value:
            flattened.extend(_flatten_strings(child))
        return flattened
    return []


def validate_report(report: InvestigationReport) -> ValidationResult:
    """Run deterministic checks that complement Pydantic's schema validation."""

    result = ValidationResult()
    grade = report.scoring.final_grade.value
    if grade not in {"A", "B", "C", "D"}:
        result.errors.append("final grade must be exactly A, B, C, or D")
    if re.fullmatch(r"[A-D][+-]", grade):
        result.errors.append("grade modifiers such as A- or B+ are prohibited")
    if not 0 <= report.scoring.total_score <= 100:
        result.errors.append("total score must be between 0 and 100")
    if report.scoring.total_score != report.scoring.calculated_total():
        result.errors.append("total score does not match scoring details")
    if report.scoring.final_grade != derive_grade(report):
        result.errors.append("grade is inconsistent with score and mandatory caps")
    if report.scoring.final_grade == Grade.A and not report.explicit_purchase_signal:
        result.errors.append("A grade requires an explicit purchase signal")
    if report.scoring.final_grade == Grade.A and report.verification.authenticity_doubtful:
        result.errors.append("an authenticity-doubtful customer cannot receive A")
    if (
        not report.customer.website
        and not report.verification.business_registry_confirmed
        and report.scoring.final_grade in {Grade.A, Grade.B}
    ):
        result.errors.append("no website plus unconfirmed entity caps the grade at C")
    if len(report.product_recommendations) > 3:
        result.errors.append("no more than three product categories may be recommended")
    if not report.priority_action.strip():
        result.errors.append("one priority action is required")
    if "\n" in report.priority_action or re.search(r"(^|\s)[12][.)]\s", report.priority_action):
        result.errors.append("priority_action must contain one action, not a list")
    if not 3 <= len(report.core_sources) <= 5:
        result.errors.append("3-5 core sources are required")
    if len({source.url for source in report.core_sources}) != len(report.core_sources):
        result.errors.append("core source URLs must be unique")
    if should_use_simplified_mode(report) and report.mode != ReportMode.SIMPLIFIED:
        result.errors.append("simplified mode trigger was ignored")

    source_ids = {source.source_id for source in report.sources}
    for item in report.investigation_items:
        missing = set(item.source_ids) - source_ids
        if missing:
            result.errors.append(f"investigation item {item.number} uses unknown sources: {missing}")
    for signal in report.purchase_signals:
        missing = set(signal.source_ids) - source_ids
        if missing:
            result.errors.append(f"purchase signal uses unknown sources: {missing}")
    for risk in report.risks:
        missing = set(risk.source_ids) - source_ids
        if missing:
            result.errors.append(f"risk uses unknown sources: {missing}")

    strings = _flatten_strings(
        report.model_dump(mode="json", exclude_computed_fields=True)
    )
    for text in strings:
        for pattern in PLACEHOLDER_PATTERNS:
            if pattern.search(text):
                result.errors.append(f"blank placeholder detected: {pattern.pattern}")
        lowered = text.lower()
        for term in FORBIDDEN_OUTPUT_TERMS:
            if term.lower() in lowered:
                result.errors.append(f"forbidden internal-process term detected: {term}")

    phone = report.customer.contact.phone
    if phone.startswith("+") and not report.customer.contact.phone.startswith("+"):
        result.errors.append("international phone prefix was removed")
    return result


def validate_rankings(entries: list[RankingEntry]) -> ValidationResult:
    result = ValidationResult()
    ranks = [entry.rank for entry in entries]
    if ranks != list(range(1, len(entries) + 1)):
        result.errors.append("batch ranking must use unique sequential ranks without ties")
    if len({entry.rank for entry in entries}) != len(entries):
        result.errors.append("batch ranking contains tied ranks")
    return result


def _docx_text(path: str | Path) -> str:
    document = Document(path)
    chunks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    return "\n".join(chunks)


def _overview_values(path: str | Path, report: InvestigationReport) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=False, read_only=False)
    sheet = workbook[SHEET_NAMES[report.report_language][0]]
    return {
        str(sheet.cell(row, 1).value): sheet.cell(row, 2).value
        for row in range(2, sheet.max_row + 1)
    }


def _expected_overview(report: InvestigationReport) -> dict[str, Any]:
    language = report.report_language
    return {
        localized(language, "公司名称", "Company name"): report.customer.legal_name,
        localized(language, "最终等级", "Final grade"): report.scoring.final_grade.value,
        localized(language, "最终得分", "Final score"): report.scoring.total_score,
        localized(language, "首推产品", "Priority products"): join_items(
            report.recommended_products, language
        ),
        localized(language, "最优先行动", "Priority action"): report.priority_action,
        localized(language, "最大风险", "Maximum risk"): report.maximum_risk,
        localized(language, "调查日期", "Investigation date"): report.investigation_date,
        localized(language, "电话", "Phone"): report.customer.contact.phone,
    }


def validate_output_consistency(
    report: InvestigationReport,
    docx_path: str | Path,
    xlsx_path: str | Path,
    json_path: str | Path,
) -> ValidationResult:
    """Ensure seven critical conclusions are identical in all three outputs."""

    result = ValidationResult()
    json_report = load_report(json_path)
    critical_fields = (
        ("customer name", report.customer.legal_name, json_report.customer.legal_name),
        ("investigation date", report.investigation_date, json_report.investigation_date),
        ("final score", report.scoring.total_score, json_report.scoring.total_score),
        ("final grade", report.scoring.final_grade, json_report.scoring.final_grade),
        ("recommended products", report.recommended_products, json_report.recommended_products),
        ("maximum risk", report.maximum_risk, json_report.maximum_risk),
        ("priority action", report.priority_action, json_report.priority_action),
    )
    for label, expected, actual in critical_fields:
        if expected != actual:
            result.errors.append(f"JSON inconsistency: {label}")

    overview = _overview_values(xlsx_path, report)
    for label, expected in _expected_overview(report).items():
        actual = overview.get(label)
        if isinstance(expected, date) and hasattr(actual, "date"):
            actual = actual.date()
        if actual != expected:
            result.errors.append(
                f"Excel inconsistency for {label!r}: expected {expected!r}, got {actual!r}"
            )

    word_text = _docx_text(docx_path)
    required_word_values = [
        report.customer.legal_name,
        report.investigation_date.isoformat(),
        str(report.scoring.total_score),
        report.scoring.final_grade.value,
        *report.recommended_products,
        report.priority_action,
    ]
    if report.maximum_risk not in word_text:
        # Ordinary reports carry risk in conclusion/nine-item findings; simplified reports
        # may carry it only in their conclusion. Requiring it in visible text catches drift.
        required_word_values.append(report.maximum_risk)
    for value in required_word_values:
        if str(value) not in word_text:
            result.errors.append(f"Word inconsistency: missing {value!r}")

    for term in FORBIDDEN_OUTPUT_TERMS:
        if term.lower() in word_text.lower():
            result.errors.append(f"Word contains forbidden internal-process term: {term}")
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input_json")
    parser.add_argument("--docx")
    parser.add_argument("--xlsx")
    parser.add_argument("--generated-json")
    args = parser.parse_args()

    report = load_report(args.input_json)
    result = validate_report(report)
    if args.docx and args.xlsx and args.generated_json:
        result.extend(
            validate_output_consistency(
                report,
                args.docx,
                args.xlsx,
                args.generated_json,
            )
        )
    for warning in result.warnings:
        print(f"WARNING: {warning}")
    for error in result.errors:
        print(f"ERROR: {error}")
    if not result.passed:
        raise SystemExit(1)
    print("PASS: report validation completed")


if __name__ == "__main__":
    main()
