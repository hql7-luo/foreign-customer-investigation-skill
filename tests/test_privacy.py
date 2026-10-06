from __future__ import annotations

import re
import subprocess
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


EMAIL = re.compile(r"\b[A-Z0-9._%+-]+@([A-Z0-9.-]+\.[A-Z]{2,})\b", re.IGNORECASE)
PHONE = re.compile(r"\+[0-9][0-9 ()-]{7,}[0-9]")
SECRET_PATTERNS = [
    re.compile(r"ghp_[A-Za-z0-9]{20,}"),
    re.compile(r"github_pat_[A-Za-z0-9_]{20,}"),
    re.compile(r"sk-[A-Za-z0-9]{20,}"),
    re.compile(r"AKIA[0-9A-Z]{16}"),
    re.compile(r"-----BEGIN (?:RSA |EC |OPENSSH )?PRIVATE KEY-----"),
]


def _public_files(root: Path) -> list[Path]:
    """Audit publication content, including new source files, without local runtime data."""

    if (root / ".git").exists():
        result = subprocess.run(
            ["git", "ls-files", "--cached", "--others", "--exclude-standard", "-z"],
            cwd=root,
            check=True,
            capture_output=True,
        )
        candidates = [root / name for name in result.stdout.decode().split("\0") if name]
    else:
        # Exported archives have no Git index; scan the packaged source/example trees.
        candidates = [path for path in root.iterdir() if path.is_file()]
        for directory in (root / ".agents", root / "examples", root / "tests"):
            candidates.extend(directory.rglob("*"))
    return sorted({
        path for path in candidates
        if path.is_file() and "__pycache__" not in path.parts
        and path.suffix.lower() in {
            ".py", ".md", ".json", ".txt", ".toml", ".yaml", ".yml", ".docx", ".xlsx", ""
        }
    })


def _artifact_text(path: Path) -> str:
    if path.suffix.lower() == ".docx":
        document = Document(path)
        chunks = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                chunks.extend(cell.text for cell in row.cells)
        return "\n".join(chunks)
    if path.suffix.lower() == ".xlsx":
        workbook = load_workbook(path, data_only=False, read_only=False)
        return "\n".join(
            str(cell.value)
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if cell.value is not None
        )
    return path.read_text(encoding="utf-8", errors="ignore")


def test_public_repository_contains_only_fictional_contact_data():
    root = Path(__file__).resolve().parents[1]
    files = _public_files(root)
    corpus = "\n".join(_artifact_text(path) for path in files)

    for pattern in SECRET_PATTERNS:
        assert not pattern.search(corpus), f"possible secret detected: {pattern.pattern}"

    domains = {match.group(1).lower() for match in EMAIL.finditer(corpus)}
    assert all(domain.endswith(".example") or domain == "example.com" for domain in domains)

    phones = {match.group(0) for match in PHONE.finditer(corpus)}
    assert all(phone.startswith("+999") for phone in phones)
    assert not any(path.name == "客户背景调查提示词.docx" for path in files)


def test_every_binary_example_is_visibly_fictional():
    root = Path(__file__).resolve().parents[1]
    artifacts = [
        path
        for base in (
            root / "examples",
            root / ".agents" / "skills" / "foreign-customer-investigation" / "templates",
        )
        for path in base.rglob("*")
        if path.suffix.lower() in {".docx", ".xlsx"}
    ]
    assert artifacts
    for artifact in artifacts:
        assert "FICTIONAL EXAMPLE — NOT A REAL COMPANY" in _artifact_text(artifact)


def test_privacy_audit_scans_publication_files_but_ignores_local_environments(tmp_path):
    subprocess.run(["git", "init", "--quiet", str(tmp_path)], check=True)
    (tmp_path / ".gitignore").write_text(".venv/\noutputs/\n", encoding="utf-8")
    tracked = tmp_path / "README.md"
    tracked.write_text("FICTIONAL public report", encoding="utf-8")
    subprocess.run(["git", "add", "README.md", ".gitignore"], cwd=tmp_path, check=True)
    source = tmp_path / "new-source.py"
    source.write_text("# Newly added publication source", encoding="utf-8")
    for directory in (".venv", "outputs"):
        (tmp_path / directory).mkdir()
        (tmp_path / directory / "local.txt").write_text("Local material", encoding="utf-8")

    public_files = _public_files(tmp_path)
    assert tracked in public_files
    assert source in public_files
    assert not any(path.parent.name in {".venv", "outputs"} for path in public_files)


def test_privacy_audit_supports_exported_source_archives(tmp_path):
    (tmp_path / "README.md").write_text("FICTIONAL public report", encoding="utf-8")
    (tmp_path / "examples").mkdir()
    example = tmp_path / "examples" / "fictional.json"
    example.write_text("{}", encoding="utf-8")
    (tmp_path / ".venv").mkdir()
    (tmp_path / ".venv" / "library.txt").write_text("Local dependency", encoding="utf-8")
    assert _public_files(tmp_path) == [tmp_path / "README.md", example]
