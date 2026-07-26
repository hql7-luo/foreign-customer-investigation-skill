from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from generate_docx import generate_docx
from generate_json import generate_json, load_report
from generate_xlsx import SHEET_NAMES, generate_xlsx
from models import InvestigationReport, ReportLanguage
from validate_report import validate_output_consistency


def test_word_generation(report_a, tmp_path):
    path = generate_docx(report_a, tmp_path / "report.docx")
    document = Document(path)
    assert document.paragraphs[0].text.startswith(report_a.customer.legal_name)
    assert len(document.tables) == 2
    assert len(document.tables[0].rows) == 10


def test_simplified_word_generation(report_c, tmp_path):
    path = generate_docx(report_c, tmp_path / "simplified.docx")
    document = Document(path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "已确认信息清单" in text
    assert "九项背景调查总表" not in text


def test_excel_generation_and_formulas(report_a, tmp_path):
    path = generate_xlsx(report_a, tmp_path / "report.xlsx")
    workbook = load_workbook(path, data_only=False)
    assert workbook.sheetnames == SHEET_NAMES[report_a.report_language]
    scoring = workbook[SHEET_NAMES[report_a.report_language][2]]
    formulas = [
        cell.value
        for row in scoring.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    ]
    assert any("SUM(" in formula for formula in formulas)
    assert any("IF(" in formula for formula in formulas)
    assert scoring.auto_filter.ref
    assert scoring.freeze_panes == "A2"


def test_json_generation(report_a, tmp_path):
    path = generate_json(report_a, tmp_path / "report.json")
    raw = path.read_text(encoding="utf-8")
    assert "虚构" in raw
    assert "\\u865a" not in raw
    restored = load_report(path)
    assert restored.scoring.total_score == report_a.scoring.total_score


def test_three_output_consistency(report_a, tmp_path):
    json_path = generate_json(report_a, tmp_path / "report.json")
    docx_path = generate_docx(report_a, tmp_path / "report.docx")
    xlsx_path = generate_xlsx(report_a, tmp_path / "report.xlsx")
    result = validate_output_consistency(report_a, docx_path, xlsx_path, json_path)
    assert result.passed, result.errors


def test_non_latin_characters_survive_all_outputs(report_a, tmp_path):
    payload = report_a.model_dump(mode="json", exclude_computed_fields=True)
    payload["customer"]["original_name"] = "ООО «Вымышленная Аврора Печать»"
    payload["customer"]["original_address"] = "Вымышленная улица, 1"
    report = InvestigationReport.model_validate(payload)

    json_path = generate_json(report, tmp_path / "russian.json")
    docx_path = generate_docx(report, tmp_path / "russian.docx")
    xlsx_path = generate_xlsx(report, tmp_path / "russian.xlsx")

    assert "ООО «Вымышленная" in json_path.read_text(encoding="utf-8")
    json_restored = load_report(json_path)
    assert json_restored.customer.original_address.startswith("Вымышленная")
    assert Document(docx_path)
    overview = load_workbook(xlsx_path)[SHEET_NAMES[report.report_language][0]]
    overview_values = {
        overview.cell(row, 1).value: overview.cell(row, 2).value
        for row in range(2, overview.max_row + 1)
    }
    assert overview_values["电话"] == report.customer.contact.phone


def test_excel_phone_prefix_preserved(report_a, tmp_path):
    path = generate_xlsx(report_a, tmp_path / "phone.xlsx")
    workbook = load_workbook(path, data_only=False)
    overview = workbook[SHEET_NAMES[report_a.report_language][0]]
    values = [overview.cell(row, 2).value for row in range(2, overview.max_row + 1)]
    assert report_a.customer.contact.phone in values


def test_json_uses_english_field_names(report_b, tmp_path):
    path = generate_json(report_b, tmp_path / "english-keys.json")
    payload = json.loads(path.read_text(encoding="utf-8"))
    assert "investigation_date" in payload
    assert "customer" in payload
    assert "scoring" in payload
    assert "调查日期" not in payload


def test_generated_templates_open():
    root = Path(__file__).resolve().parents[1]
    templates = root / ".agents" / "skills" / "foreign-customer-investigation" / "templates"
    assert Document(templates / "report-template-zh.docx")
    assert Document(templates / "report-template-en.docx")
    workbook = load_workbook(templates / "investigation-template.xlsx", data_only=False)
    assert len(workbook.sheetnames) == 6
